#!/usr/bin/env python3
"""Build the Introduction to Financial Industry assignment Word file."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm, Emu

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
BLACK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
HEADER_BG = "1B3A5F"
ALT_BG = "EEF2F6"

out = Path(__file__).resolve().parent / "AQA_Intro_Industry_Assignment.docx"
doc = Document()

for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)
    s.page_width = Inches(8.27)
    s.page_height = Inches(11.69)


def set_run_font(run, size=11, bold=False, italic=False, color=BLACK, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "B0B8C1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, size=9, bold=False, color=BLACK, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    set_cell_border(cell)


def add_para(text, *, size=11, bold=False, italic=False, color=BLACK, space_after=8, space_before=0, align="left"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading_custom(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    return p


def add_subhead(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=NAVY)
    return p


def make_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=9, bold=True, color=RGBColor(255, 255, 255), align="center")
        shade_cell(table.rows[0].cells[i], HEADER_BG)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, size=8.5, bold=(c_i == 0))
            if r_i % 2 == 1:
                shade_cell(table.rows[r_i + 1].cells[c_i], ALT_BG)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def bullet(text, *, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=size)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size)
    return p


# ---------- Cover ----------
add_para("Asia Quant Academy", size=12, bold=True, color=NAVY, align="center", space_after=2)
add_para("Introduction to Financial Industry", size=20, bold=True, color=NAVY, align="center", space_after=4)
add_para("Introduction Assignment", size=14, italic=True, color=GRAY, align="center", space_after=14)
add_para("Student: Cheung Chun Yu", size=12, align="center", space_after=2)
add_para("Programme: The University of Hong Kong — BSc(QFin)", size=12, align="center", space_after=2)
add_para("Date: 17 August 2026", size=12, align="center", space_after=16)

add_para(
    "This write-up answers all five questions. Figures are dated because headcount and AUM move quickly. "
    "Where a number comes from a third-party compiler rather than the firm itself, I say so. "
    "Question 4 maps every Q3 skill first to Asia Quant Academy Lectures 1–5, then to the HKU BSc(QFin) 2024–25 syllabus.",
    size=11,
    italic=True,
    align="justify",
    space_after=10,
)

# ---------- Q1 ----------
add_heading_custom("1. Three firms in each industry category")

add_para(
    "I mixed global names with Asia-relevant ones so the list is not only US/UK. "
    "Some firms sit in more than one box (BlackRock is an asset manager that also runs mutual funds and ETFs; "
    "Jane Street is both a proprietary trader and a market maker). I put each firm in the category that best matches "
    "how the industry usually talks about it, and I flag the overlaps in Question 5.",
    align="justify",
)

make_table(
    ["Category", "Firm 1", "Firm 2", "Firm 3"],
    [
        ["a. Asset management", "BlackRock", "China Asset Management (ChinaAMC)", "J.P. Morgan Asset Management"],
        ["b. Mutual funds", "Vanguard", "Fidelity Investments", "Hang Seng Investment Management"],
        ["c. Private equity", "KKR", "Blackstone", "Hillhouse Investment"],
        ["d. HF — Long/Short", "Millennium Management", "Point72", "Marshall Wace"],
        ["e. HF — Macro", "Bridgewater Associates", "Brevan Howard", "Caxton Associates"],
        ["f. HF — Relative Value", "Citadel", "AQR Capital Management", "D.E. Shaw"],
        ["g. HF — Quantitative", "Renaissance Technologies", "Two Sigma", "Man AHL"],
        ["h. Proprietary trading", "Jane Street", "Jump Trading", "DRW"],
        ["i. Market makers", "Citadel Securities", "Virtu Financial", "Optiver"],
    ],
    col_widths=[1.7, 1.7, 1.9, 1.9],
)

add_subhead("Why these names")
bullet("Asset management. BlackRock is the world’s largest public-markets manager (iShares ETFs plus institutional separate accounts). ChinaAMC is one of the largest onshore China managers. J.P. Morgan AM is a bank-affiliated global manager covering equities, bonds, liquidity and alternatives.")
bullet("Mutual funds. Vanguard and Fidelity are classic open-end fund complexes. Hang Seng Investment Management runs SFC-authorised funds that HK retail investors actually buy (tracker funds, mixed-asset funds). A mutual fund is a product wrapper; the manager is still an asset manager. I treat “mutual funds” as firms whose core public brand is pooled regulated funds.")
bullet("Private equity. KKR and Blackstone are diversified alternatives platforms (buyout, credit, real estate). Hillhouse is a leading Asia PE/growth investor, useful because a lot of HK/China deal flow never shows up in a US LBO list.")
bullet("HF long/short. Millennium and Point72 are multi-manager “pod shops” whose largest sleeve is still fundamental equity long/short. Marshall Wace is a large European L/S house (TOPS plus discretionary pods).")
bullet("HF macro. Bridgewater’s Pure Alpha is the textbook systematic global macro book. Brevan Howard and Caxton are discretionary global macro shops that trade rates, FX, and related derivatives.")
bullet("HF relative value. Citadel’s multi-strategy platform includes large RV books (convertibles, fixed income, stat arb). AQR’s hedge-fund line is academic-factor long/short and alternative risk premia. D.E. Shaw mixes systematic RV with discretionary sleeves.")
bullet("HF quantitative. Renaissance (Medallion plus institutional funds), Two Sigma, and Man AHL are model-driven. I put them here even though AQR and D.E. Shaw are also quantitative, so the four HF types in Q2 stay distinct.")
bullet("Prop trading. Jane Street, Jump and DRW trade the firm’s own capital. They are not marketing a 2-and-20 hedge fund to outside LPs as their main business.")
bullet("Market makers. Citadel Securities, Virtu and Optiver’s advertised job is supplying two-sided liquidity in equities, options or listed derivatives, usually as a registered market maker or equivalent.")

# ---------- Q2 ----------
add_heading_custom("2. Two hedge funds in each hedge-fund type")

add_para(
    "The brief asks for employees, a rough job-mix, AUM, and a strategy description that is not copied from the firm’s homepage. "
    "Job-mix figures are approximate: LinkedIn titles are self-reported, the People tab mixes functions, and third-party scrapers "
    "(Revelio Labs) classify titles differently from how a desk would. I therefore combine (i) what the firm itself publishes, "
    "(ii) Revelio/news headcount, and (iii) the shape of titles on the firm’s LinkedIn People page as of August 2026.",
    align="justify",
)

add_subhead("2a. Long/short — Millennium Management")
add_para(
    "Employees. Millennium’s own site (cited in 2026 industry write-ups) puts headcount at about 6,800 people across 140+ locations, "
    "with 340+ investment teams. ReelFinancial, using 2025 hiring data, had 6,670. Revelio’s scrape is higher because it can pick up contractors and affiliates. I use ~6,800 as the working number.",
    align="justify",
)
add_para(
    "Job mix (LinkedIn + structure). This is not a 200-person macro shop. Most staff are not portfolio managers. "
    "On LinkedIn, the densest current-employee titles are Software Engineer / Developer, Investment Analyst, Portfolio Manager, "
    "Quantitative Researcher, and Operations / Middle Office. A useful mental split is: investing teams (PMs + analysts + pod quants) "
    "maybe a third of headcount; technology and data another third; risk, treasury, legal, HR, investor relations and office ops the rest. "
    "That split is consistent with a platform that runs hundreds of pods and bills investors pass-through costs for the infrastructure.",
    align="justify",
)
add_para(
    "AUM. InvestmentNews compiled June 2025 financial reports and listed Millennium at $77.5 billion. Later 2026 firm-site figures "
    "quoted in secondary reporting are above $90 billion. I treat $77.5 billion (June 2025) as the dated snapshot and note that the "
    "platform has kept growing.",
    align="justify",
)
add_para(
    "Strategy (not the website). Millennium is the template “pod shop.” Each team is a small business with its own P&L, a slice of "
    "firm capital (often on the order of $100–200 million to start), and a mandate — usually fundamental equity long/short, but also "
    "stat arb, merger arb, credit, FX, commodities. Industry reporting (not the firm brochure) is what matters: a 5% drawdown from "
    "allocated capital typically halves the pod’s risk; a 7.5% drawdown ends the pod. PMs keep a share of their own P&L (often cited "
    "around 15%). Investors do not pay a classic 2% management fee so much as they reimburse the cost of running the building — "
    "comp, data, tech — and then pay a performance fee. That is why headcount can look like a mini investment bank while the public "
    "product is still “one hedge fund.” A 13F from Millennium is an aggregate of hundreds of independent books, not a house view of Apple.",
    align="justify",
)

add_subhead("2a. Long/short — Point72")
add_para(
    "Employees. Point72’s own page (as of 1 July 2026, via secondary citations) lists about 3,300 employees and 200+ investing teams. "
    "Revelio Labs had 4,211 in March 2026 — again, a broader scrape. I use ~3,300 firm-reported / ~4,000 third-party.",
    align="justify",
)
add_para(
    "Job mix. LinkedIn is heavy on Investment Analyst and Portfolio Manager because of Point72 Academy, which hires graduates, trains "
    "them for about 10 months, and feeds them into pods. Cubist Systematic Strategies (the systematic affiliate) adds Quantitative "
    "Researcher and Quant Developer titles. Software engineering is large but a smaller share of the page than at Two Sigma or Jane Street. "
    "Operations, compliance and investor relations are visible because the firm still has an outside-investor franchise (unlike a pure prop shop).",
    align="justify",
)
add_para(
    "AUM. With Intelligence’s H1 2025 Billion Dollar Club had Point72 around $37 billion. The firm’s own later figure, cited as of "
    "1 July 2026, is about $58.5 billion. I report both dates rather than pretending there is one true number.",
    align="justify",
)
add_para(
    "Strategy. Same multi-manager idea as Millennium, with two differences that show up in gossip and in how they hire. First, Steve Cohen "
    "rebuilt the firm after SAC’s insider-trading settlement, so compliance and the Academy are part of the product, not a side office. "
    "Second, drawdown rules are described as negotiated per PM rather than a single hard-coded 5% / 7.5% machine. Cubist is the "
    "systematic sleeve: computer-driven strategies in equities, futures and FX, which is why a Cubist QR posting talks about order-book "
    "features and mid-frequency futures rather than “cover this sector.” The public 13F is again a pile of independent books.",
    align="justify",
)

add_subhead("2b. Macro — Bridgewater Associates")
add_para(
    "Employees. Reuters (19 December 2025), citing people familiar with the firm, put headcount at 1,200–1,300. Wikipedia still says "
    "about 1,300. Revelio Labs had 1,862 in March 2026. I use ~1,300 as the firm-adjacent figure.",
    align="justify",
)
add_para(
    "Job mix. LinkedIn is not “trader, trader, trader.” Common current titles are Investment Associate / Research Associate, "
    "Investment Engineer, Software Engineer, and a large operations/client-service layer. Bridgewater hires a lot of new graduates "
    "into a research/investment programme and is famous for high early attrition (Bloomberg has written that about a quarter of new "
    "hires leave within two years). Engineering exists to run the systematic process, not to sit on an options pit.",
    align="justify",
)
add_para(
    "AUM. InvestmentNews (June 2025 reports): $78 billion. Reuters (December 2025) cited about $92 billion as of 30 September. "
    "The firm’s SEC brochure language as of year-end 2025 still describes discretionary client assets at institutional scale. "
    "I use $78 billion as the June 2025 snapshot used for the league table, and note the later $92 billion print.",
    align="justify",
)
add_para(
    "Strategy. Ignore the “we are systematic global macro” sentence on the website. The industry distinction that actually matters is "
    "between Pure Alpha (active directional and relative bets on markets, sized by the firm’s view of the economic machine) and "
    "All Weather (risk-balanced beta: the portfolio is built so that no single economic climate — growth up, growth down, inflation "
    "up, inflation down — dominates risk). All Weather is closer to a leveraged risk-parity product than to a hedge-fund trader "
    "taking a view on the next CPI print. That is why Bridgewater can look huge in AUM while running far fewer people than Millennium: "
    "the process is centralized, not 340 pods. Daily Observations is the research memo business that sits next to the funds, not a trading desk.",
    align="justify",
)

add_subhead("2b. Macro — Brevan Howard")
add_para(
    "Employees. Public estimates cluster around 800–900 investment and support staff across London, Geneva, New York, Hong Kong, Tel Aviv and elsewhere. "
    "LinkedIn’s company page is in that band. I use ~850 as a rough figure.",
    align="justify",
)
add_para(
    "Job mix. Compared with Bridgewater, LinkedIn is more trader / portfolio manager / macro researcher and less “investment associate class.” "
    "You still see plenty of Technologists, Quant Developers and Risk, because a modern macro book is mostly listed and OTC derivatives, not cash bonds. "
    "Investor relations and fund operations are a larger share than at a closed quant shop such as Renaissance, because Brevan still raises outside capital as a flagship macro manager.",
    align="justify",
)
add_para(
    "AUM. With Intelligence’s H1 2025 European Billion Dollar Club listed Brevan Howard at about $31 billion. A parallel table had Capula near $32 billion and Brevan near $32 billion. I use ~$31–32 billion (H1 2025).",
    align="justify",
)
add_para(
    "Strategy. Discretionary global macro: rates, FX, equity indices, credit and commodities, with a historical centre of gravity in developed-market rates and a reputation for making money in crises (the 2008-era story still follows the firm). "
    "Unlike Bridgewater, the edge is supposed to sit in a small number of senior PMs reading the same macro tape differently, not in one firm-wide “economic machine.” "
    "The Hong Kong office matters for an HK student: Asia rates/FX is a real sleeve, not a marketing sticker.",
    align="justify",
)

add_subhead("2c. Relative value — Citadel")
add_para(
    "Employees. Commentary putting Citadel near 3,300 people (up ~15% over three years) is consistent with 2025–26 hiring stories. "
    "Citadel’s public site emphasises ~270 PhDs across ~60 fields rather than a raw headcount. I use ~3,300. (Citadel Securities, the market maker, is a sister firm and is not in this number.)",
    align="justify",
)
add_para(
    "Job mix. LinkedIn for Citadel LLC is one of the most engineer-heavy hedge-fund pages: Software Engineer, Quantitative Researcher, Trader, and Credit/Commodities/Equities specialists. "
    "A common industry claim is that Citadel looks more like a technology company with a risk committee than like Millennium’s federation of pods. "
    "Griffin has been public about wanting people in the office; the location footprint is narrower than Millennium’s 140 cities.",
    align="justify",
)
add_para(
    "AUM. InvestmentNews June 2025: $67.6 billion. Citadel.com as of 1 July 2026: $71 billion of investment capital. I use those two dated prints.",
    align="justify",
)
add_para(
    "Strategy. Citadel is multi-strategy, not a single RV fund. The five books the industry names are commodities, credit and convertibles, equities, global fixed income and macro, and global quantitative strategies. "
    "Relative value lives inside several of those: convertible arb, fixed-income RV, statistical arbitrage. What is different from Millennium, according to people who have worked at both, is centralisation. "
    "Risk, capital and technology sit on one platform; PM drawdowns are negotiated rather than hard-coded the same way for every pod; the firm is more willing to run concentrated, high-Sharpe internal businesses (energy trading is the usual example) instead of 340 lookalike equity pods. "
    "That is why I placed Citadel in relative value: the economic engine is capturing spreads and dislocations with a lot of leverage and a lot of engineers, not taking a Bridgewater-style growth/inflation view.",
    align="justify",
)

add_subhead("2c. Relative value — AQR Capital Management")
add_para(
    "Employees. Revelio Labs: 981 people as of March 2026. Tracxn: 923 as of 31 May 2026. I use ~950.",
    align="justify",
)
add_para(
    "Job mix. Revelio’s occupation mix is unusually explicit: Finance and Operations 51.2%, Engineering 35.8%, Sales and Marketing 13.0%. "
    "That matches LinkedIn: lots of Research / Portfolio Management / Software Engineering, plus a real client-facing and fund-ops layer because AQR sells UCITS, mutual funds and hedge funds to the same institutions. "
    "It does not look like a prop trading floor.",
    align="justify",
)
add_para(
    "AUM. Careful with this one. InvestmentNews’s hedge-fund league table (June 2025) puts AQR at $51 billion. AQR’s Form ADV-style RAUM is much larger — on the order of $300 billion — because it includes long-only and retail vehicles. "
    "For this assignment I use $51 billion as hedge-fund AUM and flag the long-only pile in Question 5.",
    align="justify",
)
add_para(
    "Strategy. AQR is “relative value” in the academic-factor sense, not in the convertible-bond-desk sense. The house view, published in papers more than in ads, is that value, momentum, carry, defensive/quality and trend exist as premia across equities, bonds, currencies and commodities; the job is to harvest them systematically, size them by risk, and not override the model because this quarter feels special. "
    "That is why 2018–2020 was so painful for the brand (value and momentum both suffered) and why the 2023–25 recovery showed up as a huge AUM rebound in the InvestmentNews 1-year change (+31%). "
    "If Citadel’s RV is a trader’s spread, AQR’s RV is a professor’s long/short factor portfolio.",
    align="justify",
)

add_subhead("2d. Quantitative — Renaissance Technologies")
add_para(
    "Employees. The firm is famously small. Public estimates remain around 300 people, mostly in East Setauket. LinkedIn is sparse on purpose. I use ~300.",
    align="justify",
)
add_para(
    "Job mix. Almost the opposite of Millennium. LinkedIn, to the extent it exists, is Research Scientist, Mathematician, Software Developer — not “sector analyst” or “IR associate.” "
    "Support functions exist but are tiny relative to AUM. There is no Academy hiring 200 undergraduates a year to cover stocks.",
    align="justify",
)
add_para(
    "AUM. InvestmentNews June 2025: $46 billion. That number is mostly the institutional funds (RIEF, RIDA and related), not Medallion. "
    "Medallion has been largely employee-and-alumni capital for years and is capacity-constrained; quoting Renaissance’s “AUM” without that split is misleading, which I return to in Question 5.",
    align="justify",
)
add_para(
    "Strategy. The non-website version is simple and still mostly true: Medallion is a high-turnover, highly leveraged, mostly statistical-arbitrage engine on liquid futures and equities, staffed by people hired as scientists rather than as investors, and it does not want more outside money. "
    "The institutional funds are a diluted, lower-turnover cousin that can take outside capital and have produced much less spectacular returns. "
    "The secrecy, the Long Island campus, and the refusal to explain signals are the strategy as much as any factor name. Compare that with Two Sigma, which publishes a tech-company career page and talks about machine learning in public.",
    align="justify",
)

add_subhead("2d. Quantitative — Two Sigma")
add_para(
    "Employees. Revelio Labs: 2,185 as of March 2026, down from a 2023 peak near 2,570. I use ~2,200.",
    align="justify",
)
add_para(
    "Job mix. LinkedIn is dominated by Software Engineer, Quantitative Researcher, Data Scientist and related tech titles. "
    "Investment “analyst covering banks” barely exists. There is a visible operations/compliance layer, but the culture the page signals is a research lab attached to a cluster, which matches the public computing footprint.",
    align="justify",
)
add_para(
    "AUM. InvestmentNews June 2025: $50.7 billion. With Intelligence tables in 2025–26 have printed higher figures (high $50 billions). I use $50.7 billion as the dated league-table number.",
    align="justify",
)
add_para(
    "Strategy. Two Sigma’s public story is “ML on alternative data.” The less glossy version from people who have left: a collection of systematic books (equity L/S, futures, event-driven/merger, volatility, some structured credit) sharing a data platform, with the usual quant problems — overfitting, crowding, and the fact that a lot of “alternative data” decays once every other quant buys the same feed. "
    "It is closer to D.E. Shaw or Man AHL than to Renaissance in how openly it recruits, and closer to Renaissance than to Point72 in that humans are not covering stocks one by one.",
    align="justify",
)

make_table(
    ["Type", "Firm", "Employees (approx.)", "AUM (dated)", "Job-mix in one line"],
    [
        ["L/S", "Millennium", "~6,800 (firm, 2026)", "$77.5bn (Jun 2025); later >$90bn", "Pods + huge tech/ops platform"],
        ["L/S", "Point72", "~3,300 firm / ~4,200 Revelio", "~$37bn (H1’25); ~$58.5bn (Jul 2026)", "Academy analysts, PMs, Cubist quants"],
        ["Macro", "Bridgewater", "~1,300 (Reuters Dec 2025)", "$78bn (Jun 2025); ~$92bn (Sep 2025)", "Research associates + engineers"],
        ["Macro", "Brevan Howard", "~850", "~$31–32bn (H1 2025)", "Macro PMs/traders + derivatives tech"],
        ["RV", "Citadel", "~3,300", "$67.6bn (Jun 2025); $71bn (Jul 2026)", "Engineers, QRs, multi-strat traders"],
        ["RV", "AQR", "~950", "$51bn HF (Jun 2025); RAUM ≫ HF", "Research + engineering + client/ops"],
        ["Quant", "Renaissance", "~300", "$46bn (Jun 2025, mostly institutional)", "Scientists/developers; almost no IR army"],
        ["Quant", "Two Sigma", "~2,200 (Revelio Mar 2026)", "$50.7bn (Jun 2025)", "SWE / QR / data science heavy"],
    ],
)

# ---------- Q3 ----------
add_heading_custom("3. Fifteen live job descriptions")

add_para(
    "I used currently open (or currently listed) postings from firm career pages and Greenhouse, accessed 17 August 2026. "
    "Links are in the tables. I then pull only programming, mathematics, and finance-product requirements — not culture slogans.",
    align="justify",
)

add_subhead("3.1 Quant developer / software engineer (5)")

make_table(
    ["#", "Firm / role", "Programming", "Mathematics", "Finance products / knowledge"],
    [
        [
            "D1",
            "Jane Street — Software Engineer, New York\nhttps://www.janestreet.com/join-jane-street/position/4274288002/",
            "Primary language is OCaml (statically typed functional). Python is used for research, visualisation and ML. No prior OCaml required; they will teach it. Open-source culture. Interview is programming, not finance.",
            "No named courses. Implicit: algorithms, type systems, ability to reason about correctness. “Top-notch programming skills” is the bar, not a PDE course.",
            "Explicitly not required. “We don’t expect experience with finance.” The work still sits next to trading systems, market data and research clusters.",
        ],
        [
            "D2",
            "Jane Street — Software Engineer Internship, New York\nhttps://www.janestreet.com/join-jane-street/position/8599644002/",
            "Learn OCaml on the job; some teams use Python. Projects are intended to reach production (trading systems through to programming-language tooling).",
            "Same as full-time: strong programmer, not a specified math major. Intellectual curiosity over credentials.",
            "Finance background not required. Mentors come from high-performance trading systems as well as PL research.",
        ],
        [
            "D3",
            "IMC — Graduate Software Engineer, Chicago (2027 start)\nhttps://www.imc.com/us/careers/jobs/4818790101",
            "Proficiency required; Java or C++ preferred. Strong algorithms and data structures. Full software-development lifecycle on the trading platform.",
            "“Strong analytical skills.” Degree in CS, engineering or related tech. No named stochastic-calculus requirement.",
            "Interest in markets is enough; previous knowledge is not required. Day job is the trading platform and algorithms used by traders and QRs.",
        ],
        [
            "D4",
            "Optiver — Graduate Software Engineer, Chicago (2027 start)\nhttps://www.optiver.com/join-us/jobs/technology/chicago/graduate-software-engineer-2027-start/",
            "Build production systems that must react at very low latency. Interdisciplinary work with traders and researchers. Onboarding is an 8-week academy. AI tooling is mentioned as part of the stack.",
            "Not specified as measure theory; the implicit bar is CS plus the ability to reason about performance and correctness under time constraints.",
            "No prior market knowledge required. The systems price and trade listed markets (the firm is a market maker in equities, ETFs, options and related products).",
        ],
        [
            "D5",
            "Optiver — Graduate Software Engineer, Austin (2027 start)\nhttps://www.optiver.com/join-us/jobs/technology/austin/graduate-software-engineer-2027-start/",
            "Same graduate SWE family as Chicago: production code that has to react at market speed; AI tooling named as part of the stack; 8-week academy. Sister postings on the same page include Software Engineer — Real-Time Pricing and Research Infrastructure (Amsterdam), which is the “quant developer” flavour: pricing engines, not CRUD apps.",
            "Not a named stochastic-calculus bar. The neighbouring “real-time pricing” role is where the maths leaks in: engineers sit on the pricing path, not only on the website.",
            "No prior product knowledge required. The book is still listed markets (equities, ETFs, options) because that is what Optiver makes.",
        ],
    ],
)

add_para("Developer pattern. Languages: C++, Java, OCaml, Python. The graduate ads do not require finance. They do require data structures, the ability to ship, and comfort sitting next to traders. Functional programming (OCaml) is a Jane Street speciality, not an industry default.", align="justify")

add_subhead("3.2 Quant researcher (5)")

make_table(
    ["#", "Firm / role", "Programming", "Mathematics", "Finance products / knowledge"],
    [
        [
            "R1",
            "Point72 / Cubist — Quantitative Researcher, London\nhttps://boards.greenhouse.io/point72/jobs/7587888002",
            "Python, R, or C/C++. pandas, scikit-learn. Ability to wrangle large raw data. End-to-end: idea → data → model → backtest → production.",
            "Applied and theoretical statistics, linear algebra, machine learning. Feature engineering at intraday to daily horizons.",
            "Global macro systematic: futures, FX; swaps, fixed income or commodities a plus. Order book, price-volume and alternative data. Mid-frequency, not HFT.",
        ],
        [
            "R2",
            "Squarepoint — Junior Quant Researcher, ML Alpha Research\nhttps://boards.greenhouse.io/embed/job_app?token=6069464",
            "At least one of Python, kdb-q (or similar). Deploy and monitor live signal models.",
            "Masters-or-higher in CS, ML/NLP, statistics, signal processing, optimisation, mathematics. Time-series methods, ML, NLP.",
            "Systematic multi-asset alpha; the posting is method-first (signals, datasets) rather than “you must already know convertible bonds.”",
        ],
        [
            "R3",
            "Optiver — Graduate Quantitative Researcher, PhD (Austin, 2027)\nhttps://www.optiver.com/join-us/jobs/quantitative-research-and-machine-learning/austin/graduate-quantitative-researcher-phd-2027-start/",
            "Implied by the stack: research code that becomes live trading. Deep learning and “AI-enabled workflows” are named.",
            "PhD in Statistics, CS, ML, Mathematics or related STEM. Statistical and stochastic models for pricing, forecasting and risk. Hypothesis testing on large-scale order-flow data.",
            "Market making: predictive models of market behaviour, pricing, execution. Live evaluation, not a paper that never trades.",
        ],
        [
            "R4",
            "Jane Street — Quantitative Researcher Internship, New York\nhttps://www.janestreet.com/join-jane-street/position/8498547002/",
            "Comfortable with Python. Access to a huge CPU/GPU cluster. Linear models through deep learning, chosen per problem.",
            "Logical and mathematical thinking; experiment design; time-series; feature engineering. Most candidates have data-science or ML experience; no named degree is required.",
            "Finance not required up front. Intern classes cover markets and the path from a signal to production. Products are whatever Jane Street trades (equities, ETFs, options, futures, bonds, crypto in various books).",
        ],
        [
            "R5",
            "Optiver — Expressions of Interest, Graduate QR 2027 (Australia/NZ rights)\nhttps://optiver.com/working-at-optiver/career-opportunities/8450287002/",
            "Any of C, C++, Python, Java, etc. Programming is used daily, not as a checkbox.",
            "Quantitative degree: maths, statistics, physics, engineering, CS, applied/quant finance, econometrics. “Mathematical precision.”",
            "Market-making desks named in the stories: Delta-1, options pricing, execution/performance research. Options theory shows up in training even when the JD does not list “must know Black–Scholes.”",
        ],
    ],
)

add_para("Researcher pattern. Python is the common language; C++ and kdb-q appear when the horizon is shorter. Maths is statistics, linear algebra, ML, time series, and (for options MM) stochastic models. Product knowledge is optional for interns and PhD grads, and specific (futures, FX, order book, options) for experienced hires such as Cubist.", align="justify")

add_subhead("3.3 Quant trader (5)")

make_table(
    ["#", "Firm / role", "Programming", "Mathematics", "Finance products / knowledge"],
    [
        [
            "T1",
            "Jane Street — Quantitative Trader Internship, New York\nhttps://www.janestreet.com/join-jane-street/position/8617344002/",
            "General programming is a plus; no specific language required. One elective implements a strategy in Python against simulated markets.",
            "“Strong quantitative thinker”; no required major. Statistical analysis, model construction, mock trading. Electives include ML and market microstructure.",
            "Identify market signals; algorithmic trading; fair value; different market structures. Prior finance not required. Training covers the role of a market maker.",
        ],
        [
            "T2",
            "IMC — Graduate Quantitative Trader, Chicago (2027 start)\nhttps://www.imc.com/us/careers/jobs/4751729101",
            "Programming is a plus: Python, MATLAB or R. Not the same bar as the SWE posting.",
            "Degree in a quantitative field (maths, physics, CS, economics). Analytical thinking; decisions under time pressure.",
            "Wide range of instruments; execution and risk of complex portfolios. Previous market knowledge is not required. Base salary listed at $250,000 plus bonus.",
        ],
        [
            "T3",
            "Optiver — Quantitative Trading Internship (Singapore 2027; programme run from Sydney or Amsterdam)\nhttps://www.optiver.com/join-us/jobs/institutional-sales-and-trading/singapore/quantitative-trading-internship-singapore-2027/",
            "Build and backtest an algorithmic strategy on historical data. Language not named; Python is the practical default in this industry.",
            "“Mathematical precision,” logical problem-solving. Lectures on options pricing, markets, trading strategies and technology.",
            "Market making. Simulated options trading on live data. Derivative theory is taught in the internship, not assumed on day one.",
        ],
        [
            "T4",
            "Optiver — Quantitative Intern, Chicago (Summer 2027; may convert to trading or research)\nhttps://www.optiver.com/join-us/jobs/institutional-sales-and-trading/chicago/quantitative-intern-summer-2027/",
            "Python or another language. Use programming on real trading problems and tools, including AI-assisted research.",
            "STEM bachelor’s or master’s. Solid foundation in mathematics, probability and statistics. Logical reasoning.",
            "Interest in financial markets and quantitative decision-making. Live market simulations. No product list as a prerequisite.",
        ],
        [
            "T5",
            "IMC — Trader Intern 2027 (Hong Kong / APAC posting)\nhttps://www.imc.com/hk/careers/jobs/4941205101",
            "Not named as a coding job. The US sister intern posting (Chicago, https://www.imc.com/us/careers/jobs/4823923101) says Python, MATLAB or R is a plus. Same firm, same seat family, local office.",
            "Science, engineering, maths, statistics, actuarial, finance or economics, Distinction-average bar. Decisions under uncertainty; stay focused under pressure.",
            "IMC trades a wide range of instruments; they say they hire both more-mathematical traders and more hands-on, higher-risk-appetite traders. Prior market knowledge is not required. This is the posting an HKU student can actually sit in the same time zone for.",
        ],
    ],
)

add_para("Trader pattern. Programming is “plus” or Python, not C++ kernel work. Maths is probability and decision-making, not a PhD in statistics (unless the seat is really a QR seat). Finance is taught: options, microstructure, portfolio risk. Almost every graduate trader ad says prior market knowledge is not required — and then the interview is still mental maths, games, and options intuition. T5 is the local APAC/HK version of the same seat.", align="justify")

add_subhead("3.4 Cross-role summary")
make_table(
    ["Requirement", "Quant developer", "Quant researcher", "Quant trader"],
    [
        ["Programming", "C++/Java/OCaml as craft; algorithms; production systems", "Python (+ C++/kdb); pandas; research → production", "Python/MATLAB/R as a plus; not the job’s identity"],
        ["Mathematics", "Discrete maths, algorithms; little named stochastic calculus", "Stats, linear algebra, ML, time series, sometimes stochastic pricing", "Probability, expected value, speed; options maths if MM"],
        ["Finance products", "Usually none on entry; systems touch whatever the desk trades", "Experienced: futures, FX, options, order book. Grads: optional", "Taught on the job: listed options, ETFs, futures, microstructure"],
    ],
)

# ---------- Q4 ----------
add_heading_custom("4. Mapping Q3 requirements to AQA lectures and HKU BSc(QFin)")

add_para(
    "The brief asks me to link Q3’s programming, mathematics and finance-product requirements to courses at my university. "
    "I use two layers, because they do different jobs. "
    "Asia Quant Academy Lectures 1–5 are the material that actually names the seats, the products and the interview tricks. "
    "HKU BSc(QFin) 2024–25 is the degree on the transcript "
    "(https://ug.hkubs.hku.hk/f/curriculum/255030/BSc%28QFin%29.SYLL.2024-25_v20240520.pdf). "
    "Lecture 1 already gives the mapping rule: four pillars of a quant are mathematics, programming, financial products and market sense. "
    "Q3 is pillars 1–3; the JDs that say “finance not required” are still testing pillar 4 in the interview.",
    align="justify",
)

add_subhead("4.0 What Lectures 1–5 actually covered")
make_table(
    ["Lecture", "Title", "What I am using it for in Q4"],
    [
        [
            "L1",
            "Getting into the Quant World\n(Lecture 1 – New.pdf)",
            "Industry map (IB vs HF vs MM vs prop), LP/GP fund structure, S&T desks by product (equities, FI, FX, commodities), market-maker bid–ask, typical firm org chart (QT/PM vs QR/QD/SWE), four pillars, interview buckets (probability / programming / finance / behavioural).",
        ],
        [
            "L2-pre",
            "Python Basic\n(Lecture 2 – (Pre) Python Basic.pptx)",
            "Language itself: types, list comprehensions, OOP, PyPI. Then the “why Python” block that developers actually need: magic methods, context managers, lambdas, decorators, and Python vs C++ vs Java (compiled / interpreted / JVM+JIT).",
        ],
        [
            "L2",
            "Selected Topics in Python\n(Lecture 2 – New.pptx)",
            "How Python really runs: AST → bytecode → PVM, GIL, threading vs multiprocessing, and why NumPy can release the GIL. This is the bridge from “I can write Python” to “I understand why IMC wants C++.”",
        ],
        [
            "L3",
            "Technical Python for Data Analysis\n(Lecture 3 pptx)",
            "The QR/QT research stack: NumPy ndarrays and vectorisation, pandas, cleaning, joins/reshapes, group-by, plotting, scikit-learn / scipy / statsmodels in the install list. Cubist’s “pandas, scikit-learn” line is this lecture.",
        ],
        [
            "L4",
            "Art of Approximation\n(lecture 4.pdf)",
            "Trader/QR interview maths that QFin does not grade: dimensional analysis, extreme-case sanity checks, Taylor approximations (multiply, 1/(1+x), √(1+x)). Finance examples: present value and the square-root market-impact law P = c·σ·√(n/v).",
        ],
        [
            "L5",
            "Numerical Methods and Linear Algebra\n(Lecture 5.pdf)",
            "How you actually compute: Riemann/midpoint integration, forward/centered differences, Ax = b (Gaussian / LU, time complexity and error), nonlinear roots (relaxation, binary search, Newton). The market-impact model from L4 is rewritten as a linear system.",
        ],
    ],
)

add_subhead("4.1 Programming (Q3a) → AQA lectures → QFin")
make_table(
    ["What the JDs ask", "AQA lecture to cite", "Closest QFin course", "How the three fit"],
    [
        [
            "Write Python daily (Jane Street QR intern “comfortable with Python”; Optiver intern; Squarepoint)",
            "L2-pre: Python basics, native types, list comprehensions, OOP, PyPI. L3: Jupyter as the research loop.",
            "COMP1117 Computer Programming (core)",
            "COMP1117 is the transcript line. L2-pre is the actual language. Without L2-pre, COMP1117 is “I have seen a for-loop.”",
        ],
        [
            "Research / data code: pandas, scikit-learn, wrangle large raw data (Cubist R1; Squarepoint R2)",
            "L3: NumPy ndarray, vectorisation, pandas, cleaning, join/combine/reshape, group-by, scipy/statsmodels/scikit-learn install. L2: NumPy releases the GIL so vectorised C loops beat Python for-loops — which is why Cubist can ask for pandas on large data.",
            "FINA2390 Financial programming and databases (elective); FINA3351 Spreadsheet financial modeling (core)",
            "FINA2390 is the QFin bridge into finance data. FINA3351 is Excel/VBA — useful for IB, not for Cubist. The JD’s pandas line is L3, not FINA3351.",
        ],
        [
            "Production systems, algorithms, C++ or Java (IMC D3; Optiver D4/D5 “react at lightning-fast speeds”)",
            "L2-pre + L2: compiled C++ vs interpreted Python vs Java bytecode+JIT. GIL: only one Python thread in the interpreter; multiprocessing circumvents it; C++ is chosen when you cannot afford that. Context managers / exceptions: how production code releases files and locks.",
            "COMP2119 Data structures and algorithms (elective). No C++/Java core.",
            "Lecture 2 explains why the JD is written that way. COMP2119 is the interview course. QFin will not teach C++. Jane Street (D1/D2) will teach OCaml; IMC will not teach C++ from zero as happily.",
        ],
        [
            "OCaml / functional style (Jane Street D1/D2)",
            "L2-pre: first-class functions, lambdas, decorators — the closest AQA gets to “functions as values.” Not OCaml.",
            "No QFin course",
            "Honest gap on both sides. Jane Street’s own JD says they will teach OCaml. I would not pretend COMP1117 covers it.",
        ],
        [
            "kdb-q / tick data (Squarepoint)",
            "L3 covers in-memory tables (pandas), not a column-store time-series DB. L2 performance notes explain why a research language and a tick DB are different tools.",
            "IIMT3601 Database management (elective); FINA2390",
            "SQL/data modelling yes. kdb no. On-the-job or self-taught.",
        ],
        [
            "ML / NLP / deep learning (Squarepoint, Optiver PhD, Jane Street QR)",
            "L3 installs scikit-learn; it does not teach deep learning. L1 interview slide lists “Probability / Statistics / ML” as a bucket, not a syllabus.",
            "FINA4350 NLP in finance (elective); FINA4359 Data analytics, QF and blockchain (elective)",
            "QFin has ML-adjacent electives. AQA has not taught a full ML lecture yet (that sits in Part 2 of the programme map on L1). PhD Optiver is above both.",
        ],
    ],
)

add_subhead("4.2 Mathematics (Q3b) → AQA lectures → QFin")
make_table(
    ["What the JDs ask", "AQA lecture to cite", "Closest QFin course", "How the three fit"],
    [
        [
            "Linear algebra (Cubist: “foundations of … linear algebra”; hedge ratios, factor models)",
            "L5 §3: write a model as Ax = b, solve with Gaussian elimination / LU, and care about time complexity and numerical error — not only “there exists an inverse.” L4 dimensional analysis of the market-impact law is turned into a linear system in L5.",
            "MATH2014 Multivariable calculus and linear algebra (core)",
            "MATH2014 is the theory. L5 is “now invert this in floating point, and know the error.” QR interviews sit closer to L5 than to a proof of rank-nullity.",
        ],
        [
            "Calculus / Taylor (every numerical method; options greeks later)",
            "L4 §3.4: all the speed approximations are (1+x)^n ≈ 1+nx. L5: Taylor remainder for integration O(h) vs midpoint O(h²), forward difference O(h) vs centered O(h²).",
            "MATH1013 University mathematics II (core); MATH2014",
            "QFin teaches the expansion. L4/L5 teach why a trader uses the first term and a QR uses the error term.",
        ],
        [
            "Probability, statistics, mental maths under speed (Optiver T3/T4; IMC T2; Jane Street T1 “quantitative thinker”)",
            "L4 is the interview lecture: extreme-case analysis in five seconds (Gaussian integral options), numerical approximations of products, reciprocals and square roots. L1 interview slide: Probability / Statistics / ML. L1 four pillars: “be good at probability and statistics, and even better, data analytics.”",
            "STAT2601 Probability and statistics I (core); STAT2602 (elective)",
            "STAT2601 is the degree requirement. L4 is what Optiver’s “mathematical precision” actually looks like in a screen. QFin does not grade speed. That gap is the lecture.",
        ],
        [
            "Root-finding, implied vol, numerical greeks (options MM training in T3; Optiver QR “stochastic models for pricing”)",
            "L5: Newton x ← x − f(x)/f'(x) with quadratic error; binary search; relaxation and when |f'(x*)|<1. Forward/centered/second differences for greeks when you cannot differentiate the pricer in closed form.",
            "FINA3350 Mathematical finance (core); MATH3906 Financial calculus as a substitute; FINA4354 Financial engineering (capstone)",
            "FINA3350/4354 give Black–Scholes. L5 gives implied-vol as “solve f(σ)=0” and greeks as finite differences. The JD never says “Newton”; the desk does.",
        ],
        [
            "Market impact / execution research (Jane Street QT microstructure elective; HFT context in L1)",
            "L4 finance example: square-root law P = c·σ·√(n/v), then “does it make sense?” by extreme cases (n→0, v→∞). L5 rewrites the same law as a linear system from dimensions.",
            "No QFin course named “market impact.” Closest: FINA2320 (portfolio), FINA4341 (risk)",
            "This is an AQA-only link. I would cite L4/L5 in an interview before I cited FINA2320.",
        ],
        [
            "Time series, experiment design, backtests (Jane Street R4; Squarepoint R2)",
            "L3: pandas time-series-friendly tables, group-by, statsmodels in the package list. L1 programme map lists Time Series Analysis and Backtesting Methods in Part 2 — not yet taught in L1–L5.",
            "ECON2280 Introductory econometrics (core); STAT4601 / ECON3283 (electives)",
            "ECON2280 is OLS. The research JD is walk-forward and multiple testing, which AQA has named but not delivered yet.",
        ],
        [
            "Stochastic processes / risk-neutral pricing (Optiver PhD R3)",
            "Not in L1–L5. L1 programme map puts Financial Mathematics in Part 2.",
            "FINA3350; MATH3603 or STAT3603 Stochastic processes (elective)",
            "Here the university is ahead of the lectures we have so far.",
        ],
    ],
)

add_subhead("4.3 Finance products and market sense (Q3c) → AQA lectures → QFin")
make_table(
    ["What the JDs ask", "AQA lecture to cite", "Closest QFin course", "How the three fit"],
    [
        [
            "Know what a hedge fund / MM / prop shop is, and which seat you are applying to",
            "L1: LP/GP structure; market maker earns the bid–ask and manages inventory; prop/HFT trade the firm’s capital; org chart with QT/PM as risk-takers (base + PnL cut) vs QR/QD/SWE on “the desk” (base + discretion of the risk-taker).",
            "FINA3325 Alternative investments (elective). Not in QFin core.",
            "Q1–Q2 of this assignment are L1 applied. QFin will not tell you that a Jane Street SWE reports next to a trader who owns the PnL.",
        ],
        [
            "Cash equities, L/S, convertibles (pod shops; Citadel credit & convertibles; IB ECM)",
            "L1 IB primary: IPO, follow-on, equity-linked (convertible). L1 secondary equities desk: cash, equity derivatives, Delta One, convertible bond trading.",
            "FINA2320 Investments; FINA1310 Corporate finance; ACCT1101",
            "L1 names the desk. FINA2320 names the portfolio maths. Convertible arb as a trade is still neither.",
        ],
        [
            "Listed options / market making (Optiver T3 lectures on options pricing; IMC T2 “complex portfolios”)",
            "L1: market maker definition; “How market makers earn money”; HK market-maker list. L1 equities desk includes equity derivatives and Delta One — the Optiver/IMC product set.",
            "FINA2322 Derivatives (core); FINA4354 Financial engineering",
            "FINA2322 is the contract. L1 is the business model (spread vs inventory). T3 then teaches pricing in-house because neither source is a trading sim.",
        ],
        [
            "Futures, FX, swaps, commodities (Cubist R1; Bridgewater / Brevan)",
            "L1 S&T map: G10 FX, EM FX, FX options; IRS, credit, TIPS; oil, gas, power, ags, metals. That is Cubist’s “futures, FX; swaps/FI/commodities a plus” almost line for line.",
            "FINA2322; FINA3323 Fixed income; ECON1220 Macro",
            "L1 is the product dictionary the JD assumes. QFin splits it across three courses and never puts them on one slide.",
        ],
        [
            "IB vs buy-side (why a “quant developer” JD can come from either)",
            "L1: IB as supermarket; sales / traders / structurers; AM vs private bank (AUM fee, not absolute return). Buy-side insider-information warning. Shared resources (legal, compliance, risk) vs the trading team.",
            "FINA3353 Regulatory, operational and valuation issues (elective)",
            "Explains Q2’s ops/compliance headcount at AQR. Not a trading edge.",
        ],
        [
            "Market microstructure, order book, latency (Jane Street T1 elective; Cubist “order book … mid frequency”)",
            "L1 programme map lists Market Microstructure and HFT in Part 2. L4/L5 market-impact law is the only microstructure maths we have so far.",
            "No dedicated QFin course",
            "Named in AQA, not taught yet, not in QFin. I flag this rather than fake a mapping.",
        ],
        [
            "“Finance knowledge not required” on graduate trader ads (IMC, Jane Street, Optiver)",
            "L1: “At least know how to trade a product and what the product is” and “Know why things happen” (market sense). Interview still has a Finance bucket. L4: you will be asked to sanity-check PV = C/(1+i)^n in five seconds anyway.",
            "FINA2322 / FINA2320 still worth taking",
            "The JD and the interview disagree. L1/L4 explain the disagreement. I extract both, as in Q3 and Q5.",
        ],
    ],
)

add_subhead("4.4 The three seats, using Lecture 1’s org chart")
add_para(
    "Lecture 1 splits the trading firm into risk-takers (quant trader / PM: base + PnL cut) and the desk around them (QR, QD, SWE: base + discretion of the risk-taker). "
    "That is why the same firm posts three JDs with three bars.",
    align="justify",
)
make_table(
    ["Seat (L1 name)", "JDs in Q3", "AQA stack I would cite", "QFin stack I would take"],
    [
        [
            "Quant developer / software engineer (non-trading, “the desk”)",
            "D1–D5 Jane Street, IMC, Optiver",
            "L2-pre + L2 (language internals, C++ vs Python, GIL) then enough L3 to talk to researchers. L1 so I know I am not the PM.",
            "COMP1117, COMP2119, FINA2390, IIMT3601. Free-elective CS if the timetable allows. Skip treating FINA3351 as the main coding course.",
        ],
        [
            "Quant researcher (non-trading, “the desk”)",
            "R1–R5 Cubist, Squarepoint, Optiver, Jane Street",
            "L3 (pandas/NumPy/sklearn) + L5 (linear systems, Newton, finite differences) + L4 (sanity-check a formula) + L1 product map so “futures/FX” is not a surprise.",
            "STAT2601/2602, MATH2014, ECON2280, STAT4601 or ECON3283, FINA3350 or MATH3906, FINA4350/4359, FINA4341.",
        ],
        [
            "Quant trader / PM (risk taker)",
            "T1–T5 Jane Street, IMC (incl. HK intern), Optiver",
            "L4 first (speed approximations, extreme cases). L1 product and MM slides. L3 Python as a plus, matching IMC/Jane Street “programming is a plus.” L5 Newton if the seat is options.",
            "FINA2322, FINA4354, FINA2320, FINA3323, STAT2601/2602, FINA2390 so the Python elective is not new.",
        ],
    ],
)

add_subhead("4.5 What I would actually do next")
add_para(
    "Capstone: FINA4341 Quantitative risk management or FINA4354 Financial engineering — both show up in the JDs; take one as capstone and the other as elective if the timetable allows.",
    align="justify",
)
add_para(
    "AQA Part 2 on the L1 programme map (algorithmic trading, microstructure, HFT, time series, ML, risk modelling, backtests) is the material the Q3 JDs still assume and that L1–L5 have only named. "
    "Until those lectures, the honest mapping is: L1–L5 cover pillars and tools; QFin covers products and stochastic maths; microstructure and production C++ remain gaps.",
    align="justify",
)

# ---------- Q5 ----------
add_heading_custom("5. Questions and doubts from Q1–Q4")

add_para("These are genuine uncertainties, not rhetorical questions. Several of them would change how I filled the tables.", align="justify")

add_subhead("On Question 1 — the categories leak")
bullet("Is Jane Street a proprietary trader, a market maker, or a quantitative hedge fund? It is all three in practice, but it is not primarily an outside-LP hedge fund. I put it in prop trading and put Optiver in market making, but a different student could swap them and still be right.")
bullet("Is Citadel a relative-value fund or a multi-strategy fund? I used RV because that is where a lot of the P&L historically sits, but the firm’s own five-platform description is broader than the assignment’s four HF types.")
bullet("Mutual funds vs asset management: Vanguard is both. I treated “mutual funds” as firms whose public product is regulated pooled funds, and “asset management” as the broader institutional/ETF/separate-account business. The assignment list treats them as separate industries; the real world does not.")
bullet("Should Man AHL sit in quantitative HF while Man Group also runs discretionary GLG? I used the AHL sleeve. League tables often print Man Group AUM, which is not the same as AHL AUM.")

add_subhead("On Question 2 — the numbers are not one number")
bullet("Which AUM? Regulatory RAUM, hedge-fund-only AUM, “investment capital,” and 13F long-market-value are four different things. AQR’s $51 billion vs ~$300 billion is the extreme case. Citadel.com’s $71 billion “investment capital” may not equal InvestmentNews’s $67.6 billion. I dated every figure rather than averaging them.")
bullet("Medallion vs institutional Renaissance: $46 billion does not tell you the fund that actually produced the famous returns. Is it fair to call Renaissance a $46 billion quant HF when the crown jewel is closed?")
bullet("LinkedIn job-mix is biased. People with “Software Engineer” in the title are more likely to keep a complete LinkedIn than a 28-year-old analyst in a pod. Operations staff may be outsourced and invisible. Revelio’s “Finance and Operations 51%” at AQR is a classifier, not a desk map. I do not trust any mix to the nearest 5%.")
bullet("Employees of Citadel vs Citadel Securities vs Citadel Europe entities: LinkedIn company pages split and merge. Millennium’s 6,800 includes a lot of people who are not in New York. Are we counting heads or FTE?")
bullet("Pass-through fees mean Millennium’s “AUM” is not comparable to Bridgewater’s in economic terms. Investors keep a smaller share of the gross pie. Should the assignment treat AUM as size of economic engine or size of reported capital?")
bullet("For Brevan Howard I could not find a 2026 firm-official headcount. ~850 is an informed estimate. I would rather leave it rough than invent a fake precise number.")

add_subhead("On Question 3 — the JDs are written to recruit, not to describe the job")
bullet("Every graduate trader ad says “finance knowledge not required,” then the interview is still options and expected value. Which requirement should I extract — the PDF or the interview?")
bullet("Quant developer vs software engineer vs quantitative engineer vs strats: Jane Street calls it Software Engineer; IMC has both Graduate SWE and Quantitative Developer. I treated them as one family. Is that what the assignment meant by “quant developer”?")
bullet("PhD Optiver QR vs BSc Jane Street QR intern: the maths bar is not the same job. Averaging them into one “researcher requirement” overstates what a QFin bachelor is being asked for.")
bullet("Intern JDs vs full-time JDs: I used intern/graduate postings because those are the ones I can actually apply to. An experienced Cubist QR posting (2–6 years, swaps/commodities) is a different species.")
bullet("Links rot. Greenhouse tokens and Jane Street position IDs will die. The assignment asked me to “find JDs”; I cannot freeze the internet. Screenshots would be more durable than URLs.")

add_subhead("On Question 4 — lectures and QFin do not cover the same thing")
bullet("The brief said “courses available in your university.” I mapped to QFin and also to AQA Lectures 1–5, because those lectures are the only place that names QT vs QR vs QD, bid–ask inventory, and the square-root impact law. If the marker only wants HKU course codes, the AQA column is extra; if the marker is this academy, the QFin column is extra. I kept both.")
bullet("Lecture 2 says Python is interpreted and GIL-bound; IMC’s JD wants C++/Java. Mapping “programming” to L2-pre + COMP1117 is true for Jane Street QR and false for IMC SWE. Which JD is Q4 supposed to follow?")
bullet("Lecture 4 is the only source I have for trader speed maths. QFin has no module for √10 ≈ 3.167. Is L4 a “course” for the purpose of this question?")
bullet("Lecture 1’s programme map lists microstructure, HFT, time series, ML and backtests in Part 2. Those are exactly the words in the Jane Street / Cubist / Squarepoint JDs, but we have not had those lectures yet. I refused to map them to L1–L5 as if we had.")
bullet("FINA3351 Spreadsheet financial modeling is core. No JD I read, and no AQA lecture, treats Excel as the main tool. Does that mean QFin’s core is closer to AM/IB (L1 asset-management slide) than to the trading-firm org chart on the same lecture?")
bullet("STAT2601 is STAT2601 in the 2024–25 PDF and SDST2601 in some 2025–26 documents. I used STAT2601.")
bullet("If “your university” includes free electives from CS, the mapping to IMC SWE gets stronger. I stayed inside QFin plus AQA so the answer matches the degree plus this programme.")

add_subhead("On the industry more generally")
bullet("Prop trading vs hedge fund vs market maker is a legal and capital-structure distinction more than a strategy distinction. The same options model can sit at Optiver (MM), Jane Street (prop) or Citadel (HF). I am still not sure the nine boxes in Question 1 are the best way to see the industry — they are just the way this assignment asked me to cut it.")
bullet("For an HK student, most of these JDs are New York, Chicago, London or Amsterdam. Visa, time zone and “who actually interviews HKU” are not in the JDs and maybe matter more than whether I took FINA3350.")

# ---------- Sources ----------
add_heading_custom("Sources (accessed 17 August 2026)")
bullet("HKU BSc(QFin) syllabus 2024–25: https://ug.hkubs.hku.hk/f/curriculum/255030/BSc%28QFin%29.SYLL.2024-25_v20240520.pdf")
bullet("Asia Quant Academy Lecture 1 – New.pdf (Getting into the Quant World: institutions, four pillars, org chart, interview buckets).")
bullet("Asia Quant Academy Lecture 2 – (Pre) Python Basic.pptx and Lecture 2 – New.pptx (Python language, C++/Java vs Python, GIL, PVM).")
bullet("Asia Quant Academy Lecture 3 – Technical Python for Data Analysis.pptx (NumPy, pandas, cleaning, group-by, scikit-learn/scipy/statsmodels).")
bullet("Asia Quant Academy lecture 4.pdf (Art of Approximation: dimensional analysis, extreme cases, Taylor speed maths, PV and square-root market impact).")
bullet("Asia Quant Academy Lecture 5.pdf (Numerical methods and linear algebra: integration/differentiation error, Ax=b, Newton / binary search / relaxation).")
bullet("InvestmentNews, “Top 10 largest hedge funds by AUM,” 2 February 2026, figures from June 2025 reports: https://www.investmentnews.com/guides/largest-hedge-funds-by-aum/265104")
bullet("With Intelligence Billion Dollar Club materials (H1 2025 tables, including Point72, Brevan Howard).")
bullet("Reuters, “Bridgewater expands equity ownership…,” 19 December 2025 (headcount 1,200–1,300; AUM ~$92bn as of 30 Sep).")
bullet("Revelio Labs firm pages: Point72, Two Sigma, Bridgewater, AQR (headcount and, for AQR, occupation mix), March 2026.")
bullet("Citadel.com homepage (investment capital $71bn as of 1 July 2026; ~270 PhDs).")
bullet("Industry pieces on pod shops: AIStockWire “What is a pod shop?” (citing Millennium ~6,800 staff, 340+ teams; Point72 ~3,300 staff, 200+ teams, ~$58.5bn as of 1 July 2026); reporting on 5% / 7.5% Millennium drawdown rules.")
bullet("Jane Street, IMC, Optiver, Point72 Greenhouse, Squarepoint Greenhouse job postings linked in Question 3.")
bullet("LinkedIn company People tabs for the eight Q2 firms (qualitative title mix, August 2026). Self-reported and approximate.")

add_para(
    "I did not copy strategy blurbs from firm homepages or from Wikipedia as the main description, as the brief asked. League-table AUM and headcount still have to come from somewhere public; those sources are listed above.",
    size=11,
    italic=True,
    align="justify",
    space_before=8,
)

doc.save(out)
print(f"Wrote {out}")
